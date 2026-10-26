"""A walkthrough or its write-up as a Word document (.docx).

The same composers that lay out the PDF lay out the Word document, so it
has the same words, screenshots, fonts, colours and spacing. Word does the
pagination: each screenshot section is kept on one page, as in the PDF,
and headings stay with what follows them. Like the PDF it is made with
QImage and QTextDocument only, so it is safe on a worker thread.
"""

from __future__ import annotations

import io
from pathlib import Path

from PyQt6.QtCore import QBuffer, QByteArray, QIODevice, Qt, QUrl
from PyQt6.QtGui import (
    QFontInfo,
    QPageLayout,
    QPageSize,
    QTextBlock,
    QTextBlockFormat,
    QTextCharFormat,
    QTextCursor,
    QTextDocument,
    QTextFormat,
    QTextImageFormat,
    QTextTable,
)

from services.walkthrough.document import (
    _CODE_FILL,
    _FRAME_COLOR,
    _FRAME_WIDTH_PT,
    _QUOTE_GAP_PT,
    TITLE,
    DocumentFormat,
    Look,
    Walkthrough,
    _Composer,
    _Group,
    _layout_dpi,
    _MarkdownComposer,
    _text_area,
    _when,
    page_layout,
)
from services.walkthrough.storage import DOCX_NAME

#: w:pPr's children in the order the schema requires; Word refuses a file
#: with them out of order.
_PARAGRAPH_ORDER = (
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
    "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
#: Border widths are in eighths of a point.
_HAIRLINE = max(2, round(_FRAME_WIDTH_PT * 8))
_EMU_PER_PT = 12700


def write_docx(
    walkthrough: Walkthrough,
    folder: Path,
    *,
    page_size: QPageSize | None = None,
    document_format: DocumentFormat = DocumentFormat(),
) -> Path:
    """Write ``folder/walkthrough.docx``, the walkthrough as its PDF shows it.

    Raises:
        OSError: The file could not be written.
    """
    layout = page_layout(page_size)
    dpi = _layout_dpi()
    doc, groups = _Composer(_text_area(layout, dpi), dpi, document_format).layout(walkthrough)
    title = f"{document_format.title or TITLE} · {_when(walkthrough.started_at)}"
    return _write(doc, groups, folder / DOCX_NAME, layout, dpi, title)


def write_markdown_docx(
    text: str,
    path: Path,
    *,
    folder: Path,
    title: str,
    page_size: QPageSize | None = None,
    look: Look = Look(),
) -> Path:
    """Write ``text``, any Markdown, as a Word document in the walkthrough's look.

    The counterpart of ``document.write_markdown_pdf``, with the same
    arguments.

    Raises:
        OSError: The file could not be written.
    """
    layout = page_layout(page_size)
    dpi = _layout_dpi()
    doc, groups = _MarkdownComposer(_text_area(layout, dpi), dpi, folder, look).layout(text)
    return _write(doc, groups, path, layout, dpi, title)


def _write(
    doc: QTextDocument, groups: list[_Group], path: Path, layout: QPageLayout, dpi: int,
    title: str,
) -> Path:
    from docx import Document
    from docx.shared import Pt

    word = Document()
    section = word.sections[0]
    page = layout.pageSize().size(QPageSize.Unit.Point)
    section.page_width, section.page_height = Pt(page.width()), Pt(page.height())
    margins = layout.margins(QPageLayout.Unit.Point)
    section.left_margin, section.right_margin = Pt(margins.left()), Pt(margins.right())
    section.top_margin, section.bottom_margin = Pt(margins.top()), Pt(margins.bottom())
    normal = word.styles["Normal"]
    default = doc.defaultFont()
    normal.font.name = QFontInfo(default).family()
    normal.font.size = Pt(default.pointSizeF())
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    properties = word.core_properties
    properties.title = title
    # Not python-docx's name for the author.
    properties.author = properties.last_modified_by = ""
    properties.comments = "Made by OpenWhisper"
    # Each screenshot section on one page: all but its last block keep with the next.
    kept = {number for group in groups for number in range(group.first, group.last)}
    _Writer(word, doc, dpi / 72, kept).blocks(doc)
    try:
        word.save(str(path))
    except OSError as exc:
        raise OSError(f"Could not write {path}: {exc}") from exc
    return path


class _Writer:
    """Copies a QTextDocument's blocks, tables and images into a Word document.

    ``unit`` is the QTextDocument's layout units per point.
    """

    def __init__(self, word, doc: QTextDocument, unit: float, kept: set[int]) -> None:
        self.word = word
        self.doc = doc
        self.unit = unit
        self.kept = kept
        self.default_family = QFontInfo(doc.defaultFont()).family()

    def blocks(self, doc: QTextDocument) -> None:
        """Every block in order; a table's blocks as that table."""
        block = doc.begin()
        while block.isValid():
            table = QTextCursor(block).currentTable()
            if table is None:
                self._paragraph(block, self.word)
                block = block.next()
                continue
            self._table(table)
            block = table.lastCursorPosition().block().next()

    def _table(self, table: QTextTable) -> None:
        """A table ruled as in the PDF: above, under each row, firmer under the header."""
        from docx.shared import Pt

        rows, columns = table.rows(), table.columns()
        word_table = self.word.add_table(rows=rows, cols=columns)
        borders = _element("w:tblBorders")
        for side in ("top", "bottom", "insideH"):
            borders.append(_border(side, _HAIRLINE))
        word_table._tbl.tblPr.append(borders)
        for row in range(rows):
            for column in range(columns):
                cell = table.cellAt(row, column)
                word_cell = word_table.cell(row, column)
                if row == 0:
                    tcPr = word_cell._tc.get_or_add_tcPr()
                    edges = _element("w:tcBorders")
                    edges.append(_border("bottom", _HAIRLINE * 2))
                    tcPr.append(edges)
                empty = word_cell.paragraphs[0]
                block = cell.firstCursorPosition().block()
                last = cell.lastCursorPosition().block().blockNumber()
                while block.isValid() and block.blockNumber() <= last:
                    self._paragraph(block, word_cell)
                    block = block.next()
                # A new cell comes with an empty paragraph; the copied ones replace it.
                if len(word_cell.paragraphs) > 1:
                    word_cell._tc.remove(empty._p)
                for paragraph in word_cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)

    def _paragraph(self, block: QTextBlock, container) -> None:
        from docx.shared import Pt

        paragraph = container.add_paragraph()
        block_format = block.blockFormat()
        layout = paragraph.paragraph_format
        layout.space_before = Pt(block_format.topMargin() / self.unit)
        layout.space_after = Pt(block_format.bottomMargin() / self.unit)
        proportional = QTextBlockFormat.LineHeightTypes.ProportionalHeight.value
        if block_format.lineHeightType() == proportional:
            layout.line_spacing = block_format.lineHeight() / 100
        left = block_format.leftMargin() + block_format.indent() * self.doc.indentWidth()
        if left:
            layout.left_indent = Pt(left / self.unit)
        if block_format.textIndent():
            # Negative hangs a list item's marker; Word tabs to the indent after it.
            layout.first_line_indent = Pt(block_format.textIndent() / self.unit)
        layout.alignment = _alignment(block_format.alignment())
        level = block_format.headingLevel()
        layout.keep_with_next = bool(level) or block.blockNumber() in self.kept
        properties = paragraph._p.get_or_add_pPr()
        if block_format.hasProperty(QTextFormat.Property.BlockTrailingHorizontalRulerWidth):
            _add(properties, _borders(_border("bottom", _HAIRLINE)))
        elif block_format.hasProperty(QTextFormat.Property.BlockCodeFence):
            _add(properties, _shading(_CODE_FILL))
        elif block_format.hasProperty(QTextFormat.Property.BlockQuoteLevel):
            _add(properties, _borders(_border("left", _HAIRLINE, space=_QUOTE_GAP_PT)))
        if level:
            # Word's navigation pane and table of contents list it.
            outline = _element("w:outlineLvl")
            outline.set(_qn("w:val"), str(min(level, 9) - 1))
            _add(properties, outline)
        self._runs(block, paragraph)

    def _runs(self, block: QTextBlock, paragraph) -> None:
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            iterator += 1
            char_format = fragment.charFormat()
            if char_format.isImageFormat():
                self._picture(paragraph, char_format.toImageFormat())
                continue
            text = fragment.text().replace("\N{OBJECT REPLACEMENT CHARACTER}", "")
            # python-docx makes "\n" a line break and "\t" a tab.
            text = text.replace("\N{LINE SEPARATOR}", "\n")
            if not text:
                continue
            href = char_format.anchorHref()
            run = paragraph.add_run(text)
            self._format(run, char_format)
            if href:
                _hyperlink(paragraph, href).append(run._r)

    def _format(self, run, char_format: QTextCharFormat) -> None:
        from docx.shared import Pt, RGBColor

        font = char_format.font()
        if font.pointSizeF() > 0 and font.pointSizeF() != self.doc.defaultFont().pointSizeF():
            run.font.size = Pt(font.pointSizeF())
        run.font.bold = font.weight() >= 600 or None
        run.font.italic = font.italic() or None
        run.font.underline = font.underline() or None
        run.font.strike = font.strikeOut() or None
        if char_format.hasProperty(QTextFormat.Property.ForegroundBrush):
            color = char_format.foreground().color()
            run.font.color.rgb = RGBColor(color.red(), color.green(), color.blue())
        if char_format.fontFixedPitch() or char_format.fontFamilies():
            family = QFontInfo(font).family()
            if family != self.default_family:
                run.font.name = family

    def _picture(self, paragraph, image_format: QTextImageFormat) -> None:
        """The image at the size the PDF shows it, framed in the same hairline."""
        from docx.shared import Pt

        resource = QTextDocument.ResourceType.ImageResource
        image = self.doc.resource(resource.value, QUrl(image_format.name()))
        if image is None or image.isNull():
            return
        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        # PNG: lossless, as in the PDF.
        image.save(buffer, "PNG")
        buffer.close()
        shape = paragraph.add_run().add_picture(
            io.BytesIO(bytes(data)),
            width=Pt(image_format.width() / self.unit),
            height=Pt(image_format.height() / self.unit),
        )
        alt = image_format.property(QTextFormat.Property.ImageAltText)
        if alt:
            shape._inline.docPr.set("descr", str(alt))
        line = _element("a:ln")
        line.set("w", str(round(_FRAME_WIDTH_PT * _EMU_PER_PT)))
        fill = _element("a:solidFill")
        color = _element("a:srgbClr")
        color.set("val", _FRAME_COLOR.lstrip("#").upper())
        fill.append(color)
        line.append(fill)
        shape._inline.graphic.graphicData.pic.spPr.append(line)


def _alignment(alignment: Qt.AlignmentFlag):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if alignment & Qt.AlignmentFlag.AlignHCenter:
        return WD_ALIGN_PARAGRAPH.CENTER
    if alignment & Qt.AlignmentFlag.AlignJustify:
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    if alignment & Qt.AlignmentFlag.AlignRight:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return None


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _element(tag: str):
    from docx.oxml import OxmlElement

    return OxmlElement(tag)


def _add(properties, element) -> None:
    """Put ``element`` into a paragraph's w:pPr where the schema wants it."""
    tags = _PARAGRAPH_ORDER
    tag = next(name for name in tags if _qn(name) == element.tag)
    properties.insert_element_before(element, *tags[tags.index(tag) + 1:])


def _border(side: str, size: int, *, space: float = 0.0):
    border = _element(f"w:{side}")
    border.set(_qn("w:val"), "single")
    border.set(_qn("w:sz"), str(size))
    border.set(_qn("w:space"), str(round(space)))
    border.set(_qn("w:color"), _FRAME_COLOR.lstrip("#").upper())
    return border


def _borders(*borders):
    element = _element("w:pBdr")
    for border in borders:
        element.append(border)
    return element


def _shading(fill: str):
    shading = _element("w:shd")
    shading.set(_qn("w:val"), "clear")
    shading.set(_qn("w:color"), "auto")
    shading.set(_qn("w:fill"), fill.lstrip("#").upper())
    return shading


def _hyperlink(paragraph, url: str):
    """A link to ``url`` at the end of ``paragraph``, to put its runs in."""
    from docx.opc.constants import RELATIONSHIP_TYPE

    link = _element("w:hyperlink")
    link.set(_qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    paragraph._p.append(link)
    return link
