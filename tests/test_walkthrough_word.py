"""A walkthrough and its write-up as Word documents, read back with python-docx."""

import os
from datetime import datetime
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from PyQt6.QtGui import QColor, QImage, QPageSize

from services.settings import WalkthroughCopyFormat
from services.walkthrough.document import DocumentFormat, Look, Section, Walkthrough
from services.walkthrough.session import WalkthroughResult
from services.walkthrough.word import write_docx, write_markdown_docx

STARTED = datetime(2026, 10, 20, 14, 32, 5)
LETTER = QPageSize(QPageSize.PageSizeId.Letter)


@pytest.fixture(autouse=True)
def qt_app(_session_qt_application):
    return _session_qt_application


def _png(path: Path, width: int = 400, height: int = 240) -> Path:
    image = QImage(width, height, QImage.Format.Format_RGB32)
    image.fill(QColor("#3060c0"))
    assert image.save(str(path))
    return path


def _texts(word) -> list[str]:
    return [paragraph.text for paragraph in word.paragraphs]


def _pictures(word) -> list:
    return word.inline_shapes


def test_a_walkthrough_reads_as_its_pdf_does_with_each_section_kept_on_a_page(tmp_path):
    sections = (
        Section("Open the menu.", _png(tmp_path / "shot-1.png"), 4.0),
        Section("Pick the backend.\n\nThen save.", _png(tmp_path / "shot-2.png"), 72.0),
        Section("That's all."),
    )
    path = write_docx(Walkthrough(STARTED, 80, sections), tmp_path, page_size=LETTER)

    assert path == tmp_path / "walkthrough.docx"
    word = Document(str(path))
    texts = [text for text in _texts(word) if text]
    assert texts == [
        "Walkthrough",
        "Tue, Oct 20, 2026, 2:32 PM · 1:20 of speech · 2 screenshots",
        "Spoken notes with screenshots, transcribed by OpenWhisper. "
        "Each screenshot comes right after the words that describe it.",
        "Open the menu.", "Screenshot 1 · 0:04",
        "Pick the backend.", "Then save.", "Screenshot 2 · 1:12",
        "That's all.",
    ]
    assert len(_pictures(word)) == 2
    # Letter, inside the PDF's margins; pictures no wider than the text.
    section = word.sections[0]
    assert round(section.page_width.pt) == 612 and round(section.page_height.pt) == 792
    text_width = section.page_width - section.left_margin - section.right_margin
    assert all(picture.width <= text_width for picture in _pictures(word))
    # A section's words, screenshot and caption hold together; the caption
    # lets the next section go to another page.
    by_text = {paragraph.text: paragraph for paragraph in word.paragraphs if paragraph.text}
    assert by_text["Pick the backend."].paragraph_format.keep_with_next
    assert by_text["Then save."].paragraph_format.keep_with_next
    assert not by_text["Screenshot 2 · 1:12"].paragraph_format.keep_with_next
    assert word.core_properties.title == "Walkthrough · Tue, Oct 20, 2026, 2:32 PM"
    assert word.core_properties.author == ""


def test_the_documents_format_settings_apply_to_word_too(tmp_path):
    sections = (Section("Words.", _png(tmp_path / "shot-1.png"), 4.0),)
    no_extras = DocumentFormat(title="Bug report", details=False, about=False, captions=False)
    word = Document(str(write_docx(Walkthrough(STARTED, 5, sections), tmp_path,
                                   document_format=no_extras)))
    assert [text for text in _texts(word) if text] == ["Bug report", "Words."]
    assert len(_pictures(word)) == 1


def test_a_write_up_keeps_its_structure_links_and_look(tmp_path):
    _png(tmp_path / "shot-1.png")
    text = """# Changing the backend

Open **Settings** and follow [the guide](https://example.com/guide).

![The menu](shot-1.png)

## Steps

- Pick one
- Save it

> Restart after.

```
backend = "remote"
```

| Name | Value |
|---|---|
| speed | fast |

![Gone](missing.png)
"""
    path = write_markdown_docx(text, tmp_path / "write-up.docx", folder=tmp_path,
                               title="Changing the backend",
                               look=Look(heading_color="#5b2a86"))
    word = Document(str(path))
    paragraphs = {paragraph.text: paragraph for paragraph in word.paragraphs if paragraph.text}

    title = paragraphs["Changing the backend"]
    outline = title._p.pPr.find(qn("w:outlineLvl"))
    assert outline is not None and outline.get(qn("w:val")) == "0"
    assert title.runs[0].font.color.rgb == (0x5B, 0x2A, 0x86)
    assert title.runs[0].font.bold
    # The link is a real one, around its words.
    body = paragraphs["Open Settings and follow the guide."]
    assert any(run.bold for run in body.runs if run.text == "Settings")
    assert [link.address for link in body.hyperlinks] == ["https://example.com/guide"]
    # List markers hang in the indent, as in the PDF.
    item = paragraphs["\N{BULLET}\tPick one"]
    assert item.paragraph_format.first_line_indent < 0 < item.paragraph_format.left_indent
    assert paragraphs["Restart after."]._p.pPr.find(qn("w:pBdr")) is not None
    code = paragraphs['backend = "remote"']
    assert code._p.pPr.find(qn("w:shd")).get(qn("w:fill")) == "F5F1EB"
    assert [[cell.text for cell in row.cells] for row in word.tables[0].rows] == [
        ["Name", "Value"], ["speed", "fast"]]
    # Only the screenshot that exists; the heading keeps with what it heads.
    assert len(_pictures(word)) == 1
    assert paragraphs["Steps"].paragraph_format.keep_with_next


def test_word_documents_are_written_beside_the_pdf_and_copied_as_files(tmp_path):
    result = WalkthroughResult(folder=tmp_path, pdf=tmp_path / "write-up.pdf",
                               markdown=tmp_path / "write-up.md", clipboard_text="text",
                               history_text="", history_raw=None, shots=0, duration_s=1.0,
                               first_shot=None)
    assert result.docx == tmp_path / "write-up.docx"
    assert result.clipboard(WalkthroughCopyFormat.WORD) == ([result.docx], "")
    assert result.clipboard(WalkthroughCopyFormat.PDF) == ([result.pdf], "")
    assert result.clipboard(WalkthroughCopyFormat.BOTH) == ([result.pdf], "text")
    assert result.clipboard(WalkthroughCopyFormat.AGENT) == ([], "text")


def test_the_word_document_is_sized_in_points_like_the_pdf(tmp_path):
    word = Document(str(write_docx(Walkthrough(STARTED, 5, (Section("Words."),)), tmp_path)))
    assert word.styles["Normal"].font.size == Pt(11)
